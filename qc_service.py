from pathlib import Path
import warnings

import pandas as pd

from classlesson_qc import (
    build_department_combination_review,
    check_abnormal_attendance,
    check_abnormal_schedule_time,
    check_course_field_completeness,
    check_first_lesson_start_month,
    check_learning_device_teaching_method,
    check_lesson_content_type,
    check_schedule_audit_status,
    check_zero_student_class,
)
from data_loader import load_classlesson_excel, merge_and_deduplicate_excel
from labeling import apply_classlesson_standard_department, apply_standard_department
from pclv_qc import (
    build_department_combination_review as build_pclv_department_combination_review,
    check_abnormal_attendance as check_pclv_abnormal_attendance,
    check_abnormal_schedule_time as check_pclv_abnormal_schedule_time,
    check_schedule_audit_status as check_pclv_schedule_audit_status,
)
from qc import (
    check_audit_status,
    check_actual_start_month,
    check_attendance_card_setting,
    check_cancelled_class_status,
    check_class_and_actual_duration,
    check_class_closure,
    check_class_format_and_delivery,
    check_class_location_and_room,
    check_class_price,
    check_class_setup_date,
    check_class_type,
    check_external_class_time,
    check_learning_device_online_settings,
    check_management_project_class_name,
    check_minimum_and_opening_students,
    check_minimum_payroll_department,
    check_quarter_start_date,
    check_renewal_type,
    check_textbook_distribution,
    check_teacher_and_related_class,
    normalize_qc_month,
)


CLASSINFO_RULE_COLUMNS = [
    '续班类型异常',
    '班级标价规范',
    '未设置主带课教师',
    '关联班号异常',
    '审核状态异常',
    '季度与开课日期不符',
    '上课形式&授课方式&授课渠道异常',
    '取消班状态异常',
    '设班时间晚于开课时间',
    '实际开课日期不早于开课日期当月',
    '上课时间（外）未填写',
    '教材发放形式填写错误',
    '封班检查',
    '最小发薪部门为空',
    '班级时长&实际时长异常',
    '未设置上课地点/上课教室',
    '一卡通设置错误',
    '学习机线上班级设置规范',
    '开班保底人数',
    '班级类型错误',
    '疑似非本管理项目下班级',
]

CLASSLESSON_RULE_COLUMNS = [
    '异常时间排课',
    '教师考勤异常',
    '学员考勤异常',
    '0人班处理',
    '配课表审核时效异常',
    '授课内容类型异常',
    '课程字段完整性异常',
    '学习机-授课方式标识异常',
    '上课日期不早于开课日期当月',
]

PCLV_RULE_COLUMNS = [
    '异常时间排课',
    '教师考勤异常',
    '学员考勤异常',
    '配课表审核时效异常',
]

CLASSINFO_SHEET_NAMES = [
    '本月&次月开课',
    '本月设班',
    '本月结课',
]

ALL_DEPARTMENTS = '全部部门'
BUSINESS_CLASS = 'class'
BUSINESS_PREMIUM = 'premium'
PREMIUM_DEPARTMENT = '高中一对一部'


def combine_classinfo_files(file_paths, output_path):
    """
    将三份独立的 classinfo 文件合并为固定 Sheet 名的工作簿。

    每个文件只读取第一个 Sheet，并要求三份文件的列名和顺序完全一致。
    """
    if len(file_paths) != len(CLASSINFO_SHEET_NAMES):
        raise ValueError('班级信息表需要上传三份文件。')

    frames = []
    expected_columns = None
    with warnings.catch_warnings():
        warnings.filterwarnings(
            'ignore',
            message='Workbook contains no default style.*',
            category=UserWarning,
            module='openpyxl.styles.stylesheet',
        )
        for sheet_name, file_path in zip(CLASSINFO_SHEET_NAMES, file_paths):
            frame = _read_classinfo_source(file_path)
            columns = list(frame.columns)
            if expected_columns is None:
                expected_columns = columns
            elif columns != expected_columns:
                raise ValueError(
                    f'“{sheet_name}”的列结构与其他班级信息表不一致，请检查是否上传了正确文件。'
                )
            frames.append(frame)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        for sheet_name, frame in zip(CLASSINFO_SHEET_NAMES, frames):
            frame.to_excel(writer, sheet_name=sheet_name, index=False)
    return output_path


def _read_classinfo_source(file_path):
    file_path = Path(file_path)
    if file_path.suffix.lower() == '.csv':
        last_error = None
        for encoding in ['utf-8-sig', 'gb18030']:
            try:
                return pd.read_csv(file_path, encoding=encoding)
            except UnicodeDecodeError as exc:
                last_error = exc
        raise ValueError('CSV 文件编码无法识别，请使用 UTF-8 或 GB18030 编码。') from last_error
    return pd.read_excel(file_path, sheet_name=0, header=0)


def run_classinfo_qc(input_path, qc_month=None):
    result = merge_and_deduplicate_excel(str(input_path))
    result = apply_standard_department(result)

    rules = [
        check_renewal_type,
        check_class_price,
        check_teacher_and_related_class,
        check_audit_status,
        check_quarter_start_date,
        check_class_format_and_delivery,
        check_cancelled_class_status,
        check_class_setup_date,
        check_actual_start_month,
        check_external_class_time,
        check_textbook_distribution,
        check_class_closure,
        check_minimum_payroll_department,
        check_class_and_actual_duration,
        check_class_location_and_room,
        check_attendance_card_setting,
        check_learning_device_online_settings,
        check_minimum_and_opening_students,
        check_class_type,
        check_management_project_class_name,
    ]
    for rule in rules:
        if rule in {check_teacher_and_related_class, check_class_closure}:
            result = rule(result, qc_month=qc_month)
        else:
            result = rule(result)
    return result


def run_classlesson_qc(input_path):
    result = load_classlesson_excel(str(input_path))
    result = apply_classlesson_standard_department(result)
    result = check_abnormal_schedule_time(result)
    result = check_abnormal_attendance(result)
    result = check_zero_student_class(result)
    result = check_schedule_audit_status(result)
    result = check_lesson_content_type(result)
    result = check_course_field_completeness(result)
    result = check_learning_device_teaching_method(result)
    result = check_first_lesson_start_month(result)
    return result


def run_pclv_qc(input_path):
    result = load_classlesson_excel(str(input_path))
    result = _apply_pclv_standard_department(result)
    result = check_pclv_abnormal_schedule_time(result)
    result = check_pclv_abnormal_attendance(result)
    result = check_pclv_schedule_audit_status(result)
    return result


def detect_schedule_type(input_path):
    columns = _read_first_sheet_columns(input_path)
    if {'班级编号', '课次', '日期', '教师打卡状态'}.issubset(columns):
        return 'classlesson'
    if {'课时包编号', '课次序号', '学员姓名', '教师考勤'}.issubset(columns):
        return 'pclv'
    raise ValueError('无法识别配课表类型，请上传 classlesson 或 pclv 文件。')


def validate_classinfo(input_path):
    columns = _read_first_sheet_columns(input_path)
    required = {'标准部门', '班级编码', '班级名称（外）', '产品体系'}
    missing = sorted(required - columns)
    if missing:
        raise ValueError(f"班级信息表缺少字段：{', '.join(missing)}")


def run_qc_job(
    classinfo_path,
    schedule_path,
    output_path,
    qc_month=None,
    standard_department=None,
    business_type=None,
):
    schedule_type = detect_schedule_type(schedule_path)
    qc_month = normalize_qc_month(qc_month)
    selected_department = normalize_department_filter(standard_department)
    business_type = normalize_business_type(business_type, schedule_type)
    validate_business_selection(
        business_type,
        schedule_type,
        selected_department,
    )

    if schedule_type == 'classlesson':
        schedule_result = run_classlesson_qc(schedule_path)
        combination_review_builder = build_department_combination_review
        schedule_rules = CLASSLESSON_RULE_COLUMNS
        schedule_sheet_name = '配课质检-classlesson'
    else:
        schedule_result = run_pclv_qc(schedule_path)
        combination_review_builder = build_pclv_department_combination_review
        schedule_rules = PCLV_RULE_COLUMNS
        schedule_sheet_name = '配课质检-pclv'

    available_schedule_departments = list_available_departments(schedule_result)
    if (
        selected_department != ALL_DEPARTMENTS
        and selected_department not in available_schedule_departments
    ):
        raise ValueError(
            f'所选部门“{selected_department}”在配课表中不存在。'
            '请检查部门选择或上传文件。'
        )
    schedule_result = filter_by_standard_department(
        schedule_result,
        selected_department,
    )

    classinfo_checked = classinfo_required_for_selection(business_type)
    if classinfo_checked:
        if classinfo_path is None:
            raise ValueError('当前业务和部门需要上传三份班级信息表。')
        validate_classinfo(classinfo_path)
        classinfo_result = run_classinfo_qc(classinfo_path, qc_month=qc_month)
        available_classinfo_departments = list_available_departments(
            classinfo_result,
        )
        if (
            selected_department != ALL_DEPARTMENTS
            and selected_department not in available_classinfo_departments
        ):
            raise ValueError(
                f'所选部门“{selected_department}”在班级信息中不存在。'
                '请检查部门选择或上传文件。'
            )
        classinfo_result = filter_by_standard_department(
            classinfo_result,
            selected_department,
        )
    else:
        classinfo_result = pd.DataFrame(columns=['标化部门'])

    department_combination_review = combination_review_builder(schedule_result)

    classinfo_stats = (
        summarize_rules(classinfo_result, CLASSINFO_RULE_COLUMNS)
        if classinfo_checked
        else []
    )
    schedule_stats = summarize_rules(schedule_result, schedule_rules)
    department_summary = build_department_summary(
        classinfo_result,
        schedule_result,
        CLASSINFO_RULE_COLUMNS,
        schedule_rules,
    )
    classinfo_export = (
        add_qc_feedback(classinfo_result, CLASSINFO_RULE_COLUMNS)
        if classinfo_checked
        else None
    )
    schedule_export = add_qc_feedback(
        schedule_result,
        schedule_rules,
    )

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    qc_parameters = pd.DataFrame([
        {
            '参数': '业务类型',
            '值': '班课业务' if business_type == BUSINESS_CLASS else '高端业务',
            '说明': '决定配课表类型及是否执行班级信息质检',
        },
        {
            '参数': '质检月份',
            '值': qc_month,
            '说明': '0人班封班检查按该月份的第一天和最后一天判断',
        },
        {
            '参数': '质检部门',
            '值': selected_department,
            '说明': (
                '同时用于班级信息和配课表筛选'
                if classinfo_checked
                else '用于 PCLV 配课表筛选；高端业务不检查班级信息'
            ),
        },
    ])
    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        qc_parameters.to_excel(writer, sheet_name='质检参数', index=False)
        department_summary.to_excel(writer, sheet_name='部门汇总', index=False)
        if classinfo_export is not None:
            classinfo_export.to_excel(
                writer,
                sheet_name='班级信息质检',
                index=False,
            )
        schedule_export.to_excel(writer, sheet_name=schedule_sheet_name, index=False)
        if department_combination_review is not None:
            department_combination_review.to_excel(
                writer,
                sheet_name='部门组合核对',
                index=False,
            )
        _format_output_workbook(writer.book)

    return {
        'business_type': business_type,
        'business_label': (
            '班课业务' if business_type == BUSINESS_CLASS else '高端业务'
        ),
        'schedule_type': schedule_type,
        'qc_month': qc_month,
        'selected_department': selected_department,
        'selected_schedule_department': selected_department,
        'selected_classinfo_department': (
            selected_department if classinfo_checked else None
        ),
        'classinfo_checked': classinfo_checked,
        'classinfo_rows': len(classinfo_result),
        'schedule_rows': len(schedule_result),
        'classinfo_abnormal_rows': int(
            abnormal_mask(classinfo_result, CLASSINFO_RULE_COLUMNS).sum()
        ),
        'schedule_abnormal_rows': int(
            abnormal_mask(schedule_result, schedule_rules).sum()
        ),
        'department_combination_count': (
            len(department_combination_review)
            if department_combination_review is not None
            else 0
        ),
        'department_summary': department_summary.fillna('').to_dict('records'),
        'rule_stats': classinfo_stats + schedule_stats,
        'output_path': str(output_path),
    }


def normalize_business_type(business_type=None, schedule_type=None):
    """规范业务类型；旧调用未传参数时按配课表类型兼容推断。"""
    if business_type is None or not str(business_type).strip():
        return BUSINESS_PREMIUM if schedule_type == 'pclv' else BUSINESS_CLASS
    value = str(business_type).strip().lower()
    if value not in {BUSINESS_CLASS, BUSINESS_PREMIUM}:
        raise ValueError('业务类型不正确，请选择班课业务或高端业务。')
    return value


def validate_business_selection(
    business_type,
    schedule_type,
    selected_schedule_department,
):
    """强制校验业务、部门和上传配课表类型之间的一致性。"""
    if business_type == BUSINESS_CLASS and schedule_type != 'classlesson':
        raise ValueError('班课业务只能上传班级配课表（classlesson）。')
    if business_type == BUSINESS_PREMIUM and schedule_type != 'pclv':
        if selected_schedule_department == PREMIUM_DEPARTMENT:
            raise ValueError('高中一对一部只能上传 PCLV 高端配课表。')
        raise ValueError('高端业务只能上传 PCLV 高端配课表。')
    if (
        selected_schedule_department == PREMIUM_DEPARTMENT
        and business_type != BUSINESS_PREMIUM
    ):
        raise ValueError('高中一对一部请选择“高端业务”。')


def classinfo_required_for_selection(
    business_type,
    selected_schedule_department=None,
):
    """班课业务检查班级信息；高端业务只检查 PCLV 配课表。"""
    return business_type == BUSINESS_CLASS


def normalize_department_filter(standard_department=None):
    """规范部门筛选参数；未提供时兼容为查看全部部门。"""
    if standard_department is None:
        return ALL_DEPARTMENTS
    value = str(standard_department).strip()
    return value or ALL_DEPARTMENTS


def _standard_department_series(df):
    return (
        df['标化部门']
        .astype('string')
        .str.strip()
        .fillna('未标化')
        .replace('', '未标化')
    )


def list_available_departments(*frames):
    """返回本次质检数据中实际存在的标化部门。"""
    departments = set()
    for frame in frames:
        departments.update(_standard_department_series(frame).tolist())
    return sorted(departments)


def filter_by_standard_department(df, standard_department):
    """按标化部门过滤明细；“全部部门”保留所有数据。"""
    selected_department = normalize_department_filter(standard_department)
    if selected_department == ALL_DEPARTMENTS:
        return df.copy()
    return df.loc[
        _standard_department_series(df).eq(selected_department)
    ].copy()


def abnormal_mask(df, rule_columns):
    mask = pd.Series(False, index=df.index)
    for column in rule_columns:
        if column not in df.columns:
            continue
        values = df[column]
        if column == '封班检查':
            text_values = values.astype('string').str.strip()
            mask |= text_values.notna() & text_values.ne('')
        else:
            mask |= pd.to_numeric(values, errors='coerce').eq(1)
    return mask


def add_qc_feedback(df, rule_columns):
    """在首列汇总每行命中的全部异常，多个异常使用“+”连接。"""
    result = df.copy()
    if '质检结果反馈' in result.columns:
        result = result.drop(columns=['质检结果反馈'])

    feedback_parts = [[] for _ in range(len(result))]
    for column in rule_columns:
        if column not in result.columns:
            continue

        if column == '封班检查':
            values = result[column].astype('string').str.strip()
            for position, value in enumerate(values.tolist()):
                if pd.notna(value) and value:
                    feedback_parts[position].append(str(value))
        else:
            abnormal = pd.to_numeric(
                result[column],
                errors='coerce',
            ).eq(1)
            for position, is_abnormal in enumerate(abnormal.tolist()):
                if is_abnormal:
                    feedback_parts[position].append(column)

    result.insert(
        0,
        '质检结果反馈',
        ['+'.join(parts) for parts in feedback_parts],
    )
    return result


def summarize_rules(df, rule_columns):
    stats = []
    for column in rule_columns:
        if column not in df.columns:
            continue
        if column == '封班检查':
            values = df[column].astype('string').str.strip()
            for label, count in values.value_counts().items():
                if pd.notna(label) and label:
                    stats.append({'rule': label, 'count': int(count)})
        else:
            count = pd.to_numeric(df[column], errors='coerce').eq(1).sum()
            stats.append({'rule': column, 'count': int(count)})
    return stats


def build_department_summary(
    classinfo_result,
    schedule_result,
    classinfo_rules,
    schedule_rules,
):
    classinfo_summary = _summarize_by_department(
        classinfo_result,
        classinfo_rules,
        '班级信息',
    )
    schedule_summary = _summarize_by_department(
        schedule_result,
        schedule_rules,
        '配课',
    )
    return pd.concat(
        [classinfo_summary, schedule_summary],
        ignore_index=True,
    ).sort_values(['数据类型', '标化部门'], kind='stable')


def _summarize_by_department(df, rule_columns, data_type):
    department = df['标化部门'].astype('string').fillna('未标化').replace('', '未标化')
    working = pd.DataFrame({
        '标化部门': department,
        '异常记录': abnormal_mask(df, rule_columns).astype(int),
    })
    summary = working.groupby('标化部门', dropna=False).agg(
        总记录数=('异常记录', 'size'),
        异常记录数=('异常记录', 'sum'),
    ).reset_index()
    summary.insert(0, '数据类型', data_type)
    summary['异常率'] = (
        summary['异常记录数'] / summary['总记录数']
    ).round(4)
    return summary


def _apply_pclv_standard_department(df):
    mapped = df.copy()
    department = mapped['标准部门名称'].astype('string').str.strip()
    # PCLV 不做班课式部门映射，标化部门直接沿用“标准部门名称”。
    mapped['标化部门'] = department.fillna('')
    return mapped[['标化部门'] + [c for c in mapped.columns if c != '标化部门']]


def _read_first_sheet_columns(input_path):
    with warnings.catch_warnings():
        warnings.filterwarnings(
            'ignore',
            message='Workbook contains no default style.*',
            category=UserWarning,
            module='openpyxl.styles.stylesheet',
        )
        frame = pd.read_excel(input_path, sheet_name=0, nrows=0)
    return set(frame.columns)


def _format_output_workbook(workbook):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    header_fill = PatternFill('solid', fgColor='183B4E')
    header_font = Font(color='FFFFFF', bold=True)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
        sheet.row_dimensions[1].height = 24
        for column_cells in sheet.iter_cols(min_row=1, max_row=min(sheet.max_row, 200)):
            width = min(max(len(str(cell.value or '')) for cell in column_cells) + 2, 32)
            sheet.column_dimensions[column_cells[0].column_letter].width = max(width, 10)
    if '部门汇总' in workbook.sheetnames:
        summary = workbook['部门汇总']
        for cell in summary['E'][1:]:
            cell.number_format = '0.0%'

    feedback_fill = PatternFill('solid', fgColor='FFFFFF00')
    feedback_font = Font(color='FFFF0000', bold=True)
    feedback_sheets = [
        name
        for name in workbook.sheetnames
        if name == '班级信息质检' or name.startswith('配课质检-')
    ]
    for sheet_name in feedback_sheets:
        sheet = workbook[sheet_name]
        if sheet['A1'].value != '质检结果反馈':
            continue
        sheet.column_dimensions['A'].width = 48
        for row_index, cell in enumerate(sheet['A'], start=1):
            cell.fill = feedback_fill
            cell.font = feedback_font
            cell.alignment = Alignment(
                horizontal='center' if row_index == 1 else 'left',
                vertical='center',
                wrap_text=True,
            )
    if '部门组合核对' in workbook.sheetnames:
        review = workbook['部门组合核对']
        if review.max_row >= 2:
            conclusion_validation = DataValidation(
                type='list',
                formula1='"待核对,匹配,不匹配"',
                allow_blank=False,
            )
            conclusion_validation.promptTitle = '部门组合核对'
            conclusion_validation.prompt = '请选择待核对、匹配或不匹配。'
            conclusion_validation.errorTitle = '核对结论无效'
            conclusion_validation.error = '请从下拉列表中选择核对结论。'
            conclusion_validation.showInputMessage = True
            conclusion_validation.showErrorMessage = True
            review.add_data_validation(conclusion_validation)
            conclusion_validation.add(f'E2:E{review.max_row}')
